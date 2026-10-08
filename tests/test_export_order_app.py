"""Actual SQLite metadata -> Streamlit preview -> both downloaded formats."""

import copy
import csv
import io
import os
import json
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook
from streamlit.delta_generator import DeltaGenerator
from price_test_fixtures import capture_downloads

PRICE_HEADERS = ['단가(원)', '합계(원)']
PRICE_VALUES = [None, None]

import exports
import rarities
import recognition
import references
from export_order_fixtures import make_database, monster, spell_trap
from test_app import download_buttons_disabled, row, start_app


class ExportOrderAppTest(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.database = Path(directory.name) / "inventory.sqlite"
        make_database(self.database, [
            (1, monster("通常", level=1), 90001),
            (2, monster(level=7), 90002),
            (3, spell_trap("魔法"), 90003),
            (4, spell_trap("罠"), 90004),
            (5, monster("融合／ペンデュラム／効果"), 90005),
            (6, monster("リンク／効果", level=4), 90006),
        ])
        self.files = {}
        download_patch = mock.patch.object(DeltaGenerator, 'download_button', capture_downloads(self.files))
        download_patch.start()
        self.addCleanup(download_patch.stop)
        environment = mock.patch.dict(os.environ, {'DATABASE_URL': '', 'REDIS_URL': ''})
        environment.start()
        self.addCleanup(environment.stop)
        actual_sort = exports.sort_cards

        for patch in (
            mock.patch.object(recognition, "models_ready", return_value=True),
            mock.patch.object(references, "library_info", return_value=None),
            mock.patch.object(rarities, "get_rarities", return_value=["N", "SR"]),
            mock.patch.object(exports, "sort_cards", side_effect=lambda cards:
                              actual_sort(cards, database_path=self.database)),
        ):
            patch.start()
            self.addCleanup(patch.stop)

    def prepared_app(self):
        rows = []
        for index, cid in enumerate([6, 3, 2, 4, 5, 1, 2]):
            card = {"cid": cid, "name_ko": f"카드 {cid}", "name_ja": f"カード {cid}",
                    "name_en": f"Card {cid}", "image_path": None,
                    "source_url": "https://www.db.yugioh-card.com/"}
            rows.append(row(str(index), card=card, quantity=2 if index == 6 else 1))
        app = start_app(rows).run()
        app.selectbox(key="rarity_2").set_value("SR").run()
        return app

    def test_preview_and_downloads_share_order_without_reordering_edit_rows(self):
        app = self.prepared_app()
        before = copy.deepcopy(app.session_state["rows"])
        self.assertFalse(app.exception)
        self.assertEqual(download_buttons_disabled(app), [True, True])
        expected = [
            ["카드 1", "노멀", 1], ["카드 2", "노멀", 2], ["카드 2", "슈퍼 레어", 1],
            ["카드 3", "노멀", 1], ["카드 4", "노멀", 1], ["카드 5", "노멀", 1],
            ["카드 6", "노멀", 1],
        ]
        expected = [values[:2] + ["일본판"] + values[2:] + PRICE_VALUES for values in expected]
        self.assertEqual(app.dataframe[0].value.values.tolist(), expected)
        app.checkbox(key="confirmed").check().run()
        header = ["카드명", "레어도", "판본", "수량", *PRICE_HEADERS]
        self.assertEqual(list(csv.reader(io.StringIO(self.files["csv"].decode("utf-8-sig")))),
                         [header, *[["" if v is None else str(v) for v in values] for values in expected]])
        self.assertEqual(list(load_workbook(io.BytesIO(self.files["xlsx"])).active.values),
                         [tuple(header), *[tuple(values) for values in expected]])
        self.assertEqual(app.session_state["rows"], before)
        app.radio[0].set_value("영어").run()
        self.assertFalse(app.checkbox(key="confirmed").value)
        self.assertEqual(list(app.dataframe[0].value["카드명"]),
                         ["Card 1", "Card 2", "Card 2", "Card 3", "Card 4", "Card 5", "Card 6"])

    def test_missing_passcodes_allow_preview_and_both_downloads(self):
        with closing(sqlite3.connect(self.database)) as connection, connection:
            connection.execute("UPDATE cards SET ygoprodeck_id=NULL WHERE cid IN (2, 4, 5)")
        app = self.prepared_app()
        before = copy.deepcopy(app.session_state["rows"])
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        expected = [
            ["카드 1", "노멀", 1], ["카드 2", "노멀", 2], ["카드 2", "슈퍼 레어", 1],
            ["카드 3", "노멀", 1], ["카드 4", "노멀", 1], ["카드 5", "노멀", 1],
            ["카드 6", "노멀", 1],
        ]
        expected = [values[:2] + ["일본판"] + values[2:] + PRICE_VALUES for values in expected]
        self.assertEqual(app.dataframe[0].value.values.tolist(), expected)
        app.checkbox(key="confirmed").check().run()
        self.assertFalse(app.error)
        self.assertEqual(download_buttons_disabled(app), [False, False])
        header = ["카드명", "레어도", "판본", "수량", *PRICE_HEADERS]
        self.assertEqual(list(csv.reader(io.StringIO(self.files["csv"].decode("utf-8-sig")))),
                         [header, *[["" if v is None else str(v) for v in values] for values in expected]])
        self.assertEqual(list(load_workbook(io.BytesIO(self.files["xlsx"])).active.values),
                         [tuple(header), *[tuple(values) for values in expected]])
        self.assertEqual(app.session_state["rows"], before)

    def test_metadata_failure_removes_previous_preview_blocks_downloads_and_requires_reconfirmation(self):
        app = self.prepared_app()
        app.checkbox(key="confirmed").check().run()
        self.assertEqual(download_buttons_disabled(app), [False, False])
        before = copy.deepcopy(app.session_state["rows"])
        with closing(sqlite3.connect(self.database)) as connection, connection:
            connection.execute("UPDATE cards SET list_info=? WHERE cid=5", ('{"ja": {}}',))
        app.run()
        self.assertFalse(app.exception)
        self.assertTrue(any("5" in error.value for error in app.error))
        self.assertEqual(len(app.dataframe), 0)
        self.assertEqual(download_buttons_disabled(app), [True, True])
        self.assertFalse(app.checkbox(key="confirmed").value)
        self.assertEqual(app.session_state["rows"], before)
        with closing(sqlite3.connect(self.database)) as connection, connection:
            connection.execute("UPDATE cards SET list_info=? WHERE cid=5",
                               (json.dumps({"ja": monster("融合／ペンデュラム／効果")}),))
        app.run()
        self.assertFalse(app.error)
        self.assertEqual(len(app.dataframe), 1)
        self.assertEqual(download_buttons_disabled(app), [True, True])

    def test_mapping_removed_from_tied_card_reorders_and_requires_reconfirmation(self):
        with closing(sqlite3.connect(self.database)) as connection, connection:
            connection.execute("UPDATE cards SET list_info=? WHERE cid=1",
                               (json.dumps({"ja": monster(level=7)}),))
        app = self.prepared_app()
        self.assertEqual(list(app.dataframe[0].value["카드명"])[:3], ["카드 1", "카드 2", "카드 2"])
        app.checkbox(key="confirmed").check().run()
        with closing(sqlite3.connect(self.database)) as connection, connection:
            connection.execute("UPDATE cards SET ygoprodeck_id=NULL WHERE cid=1")
        app.run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        self.assertEqual(list(app.dataframe[0].value["카드명"])[:3], ["카드 2", "카드 2", "카드 1"])
        self.assertFalse(app.checkbox(key="confirmed").value)
        self.assertEqual(download_buttons_disabled(app), [True, True])

    def test_changed_actual_order_invalidates_previous_confirmation(self):
        app = self.prepared_app()
        app.checkbox(key="confirmed").check().run()
        with closing(sqlite3.connect(self.database)) as connection, connection:
            connection.execute("UPDATE cards SET list_info=? WHERE cid=1",
                               (json.dumps({"ja": monster("効果", level=1)}),))
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(list(app.dataframe[0].value["카드명"])[:3], ["카드 2", "카드 2", "카드 1"])
        self.assertFalse(app.checkbox(key="confirmed").value)
        self.assertEqual(download_buttons_disabled(app), [True, True])


if __name__ == "__main__":
    unittest.main()
