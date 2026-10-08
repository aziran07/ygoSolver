"""Codex contract for displayed prices in files; numeric expected values are independent."""
import csv
import io
import unittest
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook
import exports

PRICE_HEADERS = ["단가(원)", "합계(원)"]
PRICE_OPTIONS = ("status", "product_url", "observed_at", "card_number")
OPTION_HEADERS = ("가격 상태", "가격 상품 URL", "가격 관찰 시각(KST)", "가격 수록 번호")
OPTION_VALUES = ("ok", "http://www.tcgshop.co.kr/goods_detail.php?goodsIdx=39577",
                 "2026-10-08T16:19:37+09:00", "15AY-JPB22")


def quote(status="ok"):
    observed = datetime(2026, 10, 8, 7, 19, 37, tzinfo=timezone.utc)
    priced = status in ('ok', 'out_of_stock', 'stock_unknown')
    return {"status": status, "detail": "수집 상품 중 확인된 재고 최저가", "locale": "ja",
            "unit_price_krw": 240 if priced else None,
            "card_number": "15AY-JPB22" if priced else None,
            "product_id": "39577" if priced else None,
            "product_url": "http://www.tcgshop.co.kr/goods_detail.php?goodsIdx=39577" if priced else None,
            "observed_at": observed if priced else None,
            "expires_at": observed + timedelta(hours=12) if priced else None,
            "matched_products": 1, "in_stock_products": 1 if status == "ok" else 0,
            "excluded_unverified": 0}


def entry(quantity=1, price=None):
    return {"cid": 5631, "name": "확산하는 파동", "rarity": "N", "quantity": quantity,
            "name_ko": "확산하는 파동", "name_ja": "拡散する波動", "name_en": "Diffusion Wave-Motion",
            "price": quote() if price is None else price}


class PriceExportTest(unittest.TestCase):
    def test_default_price_columns_are_only_numeric_unit_and_total_in_both_formats(self):
        cards = exports.aggregate([entry(2), entry(3)])
        rows = list(csv.reader(io.StringIO(exports.to_csv_bytes(cards, include_price=True).decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["카드명", "레어도", "수량", *PRICE_HEADERS])
        self.assertEqual(rows[1], ["확산하는 파동", "노멀", "5", "240", "1200"])
        sheet = load_workbook(io.BytesIO(exports.to_xlsx_bytes(cards, include_price=True))).active
        self.assertEqual(sheet["D2"].value, 240)
        self.assertEqual(sheet["E2"].value, 1200)
        self.assertEqual(sheet["D2"].data_type, "n")
        self.assertEqual(sheet.max_column, 5)

    def test_all_price_option_subsets_and_name_options_have_exact_contents(self):
        cards = exports.aggregate([entry(2), entry(3)])
        for mask in range(16):
            indices = [i for i in range(4) if mask & (1 << i)]
            # Reverse inputs to prove canonical output order rather than caller order.
            selected = tuple(PRICE_OPTIONS[i] for i in reversed(indices))
            expected_headers = ["카드명", "레어도", "수량", "공식 CID", *PRICE_HEADERS,
                                *[OPTION_HEADERS[i] for i in indices]]
            expected = ["확산하는 파동", "노멀", 5, 5631, 240, 1200,
                        *[OPTION_VALUES[i] for i in indices]]
            with self.subTest(mask=mask):
                args = {"optional_fields": ("cid",), "include_price": True, "optional_price_fields": selected}
                rows = list(csv.reader(io.StringIO(exports.to_csv_bytes(cards, **args).decode("utf-8-sig"))))
                self.assertEqual(rows, [expected_headers, [str(value) for value in expected]])
                sheet = load_workbook(io.BytesIO(exports.to_xlsx_bytes(cards, **args))).active
                self.assertEqual(list(sheet.values), [tuple(expected_headers), tuple(expected)])

    def test_nonprice_status_has_blank_amounts_and_optional_explicit_status(self):
        for status in ("not_observed", "unverified", "expired", "config_error",
                       "database_error", "cache_error", "data_error", "official_error"):
            with self.subTest(status=status):
                q = {**quote(status), "detail": "=explicit reason"}
                cards = [entry(price=q)]
                args = {"include_price": True, "optional_price_fields": PRICE_OPTIONS}
                rows = list(csv.reader(io.StringIO(exports.to_csv_bytes(cards, **args).decode("utf-8-sig"))))
                self.assertEqual(rows[1][3:], ["", "", status, "", "", ""])
                self.assertNotIn("가격 상세", rows[0])
                self.assertNotIn("가격 판본", rows[0])
                sheet = load_workbook(io.BytesIO(exports.to_xlsx_bytes(cards, **args))).active
                self.assertIsNone(sheet["D2"].value)
                self.assertIsNone(sheet["E2"].value)
                self.assertEqual(sheet["F2"].value, status)

    def test_optional_text_remains_formula_safe(self):
        cards = [entry(price={**quote(), "product_url": "=DANGEROUS()", "card_number": "+CMD"})]
        args = {"include_price": True, "optional_price_fields": ("product_url", "card_number")}
        rows = list(csv.reader(io.StringIO(exports.to_csv_bytes(cards, **args).decode("utf-8-sig"))))
        self.assertEqual(rows[1][-2:], ["'=DANGEROUS()", "'+CMD"])
        sheet = load_workbook(io.BytesIO(exports.to_xlsx_bytes(cards, **args))).active
        self.assertEqual([(sheet[cell].value, sheet[cell].data_type) for cell in ("F2", "G2")],
                         [("=DANGEROUS()", "s"), ("+CMD", "s")])

    def test_unknown_or_removed_price_options_are_rejected(self):
        for field in ("detail", "locale", "typo"):
            for enabled in (True, False):
                for writer in (exports.to_csv_bytes, exports.to_xlsx_bytes):
                    with self.subTest(field=field, enabled=enabled, writer=writer.__name__), self.assertRaises(ValueError):
                        writer([entry()], include_price=enabled, optional_price_fields=(field,))

    def test_sold_out_and_unknown_prices_are_numeric_and_keep_availability_in_both_formats(self):
        for status, label in [('out_of_stock', '품절'), ('stock_unknown', '재고 확인 불가')]:
            q = {**quote(status), 'detail': label}
            cards = exports.aggregate([entry(2, q), entry(3, q)])
            args = {"include_price": True, "optional_price_fields": PRICE_OPTIONS}
            rows = list(csv.DictReader(io.StringIO(exports.to_csv_bytes(cards, **args).decode('utf-8-sig'))))
            self.assertEqual((rows[0]['단가(원)'], rows[0]['합계(원)'], rows[0]['가격 상태']), ('240', '1200', status))
            self.assertNotIn('가격 상세', rows[0])
            self.assertEqual(rows[0]['가격 수록 번호'], '15AY-JPB22')
            self.assertTrue(rows[0]['가격 관찰 시각(KST)'])
            sheet = load_workbook(io.BytesIO(exports.to_xlsx_bytes(cards, **args))).active
            self.assertEqual((sheet['D2'].value, sheet['E2'].value, sheet['F2'].value), (240, 1200, status))
            self.assertEqual(sheet['D2'].data_type, 'n')

    def test_unavailable_stock_prices_still_require_valid_money_and_observation(self):
        for status in ('out_of_stock', 'stock_unknown'):
            for bad in ({'unit_price_krw': 0}, {'unit_price_krw': True}, {'unit_price_krw': 3.5},
                        {'unit_price_krw': None}, {'observed_at': None}, {'observed_at': datetime(2026, 10, 8)}):
                for writer in (exports.to_csv_bytes, exports.to_xlsx_bytes):
                    with self.subTest(status=status, bad=bad, writer=writer.__name__), self.assertRaises(ValueError):
                        writer([entry(price={**quote(status), **bad})], include_price=True)

    def test_aggregation_refuses_conflicting_prices_or_language(self):
        for changed in ({"unit_price_krw": 300}, {"locale": "ko"}, {"status": "expired", "unit_price_krw": None}):
            with self.subTest(changed=changed), self.assertRaises(ValueError):
                exports.aggregate([entry(), entry(price={**quote(), **changed})])

    def test_price_enabled_export_does_not_invent_missing_or_invalid_prices(self):
        missing = entry()
        del missing["price"]
        bad_quotes = [None, {}, {**quote(), "unit_price_krw": 0}, {**quote(), "unit_price_krw": True},
                      {**quote(), "unit_price_krw": 2.5}, {**quote("expired"), "unit_price_krw": 240}]
        for card in [missing, *[{**entry(), "price": q} for q in bad_quotes]]:
            for writer in (exports.to_csv_bytes, exports.to_xlsx_bytes):
                with self.subTest(card=card, writer=writer.__name__), self.assertRaises(ValueError):
                    writer([card], include_price=True)

    def test_aggregation_cannot_assign_one_rows_price_to_an_unpriced_copy(self):
        priced = entry()
        unpriced = entry()
        del unpriced["price"]
        for entries in ([priced, unpriced], [unpriced, priced]):
            with self.subTest(order=["price" in e for e in entries]), self.assertRaises(ValueError):
                exports.aggregate(entries)


if __name__ == "__main__":
    unittest.main()
