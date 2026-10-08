import csv
import io
import unittest
from fractions import Fraction
from itertools import combinations

from openpyxl import load_workbook

import exports


def entry(cid, name, quantity=1, rarity="N"):
    return {"cid": cid, "name": name, "name_ko": "한국어 " + name, "name_ja": "日本語" + name,
            "name_en": name, "quantity": quantity, "rarity": rarity}


class QuantityTest(unittest.TestCase):
    def test_rejects_non_positive_and_non_integral(self):
        for bad in [0, -1, True, False, 1.0, 2.5, Fraction(2, 1), "3", None]:
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                exports.validate_quantity(bad)

    def test_accepts_positive_int(self):
        self.assertEqual(exports.validate_quantity(3), 3)


class AggregateTest(unittest.TestCase):
    def test_distinct_rarities_split_same_cid_and_matching_rarities_merge(self):
        cards = exports.aggregate([entry(1, "Ash", 2, "SR"), entry(1, "Ash", 3, "N"),
                                   entry(1, "Ash", 4, "SR")])
        self.assertEqual([(c["cid"], c["rarity"], c["quantity"]) for c in cards],
                         [(1, "SR", 6), (1, "N", 3)])

    def test_missing_blank_unknown_and_tcg_rarities_are_rejected(self):
        missing = entry(1, "Ash")
        del missing["rarity"]
        for bad in [missing, *[entry(1, "Ash", rarity=r) for r in (None, "", "???", "STARLIGHT", 1)]]:
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                exports.aggregate([bad])

    def test_conflicting_names_rejected_even_with_different_rarities(self):
        with self.assertRaisesRegex(ValueError, "CID 1"):
            exports.aggregate([entry(1, "Ash", rarity="N"), entry(1, "Different", rarity="SR")])

    def test_same_code_color_variants_remain_distinct_in_export(self):
        cards = exports.aggregate([entry(1, "Ash", rarity="SE"), entry(1, "Ash", rarity="SE@43"),
                                   entry(1, "Ash", 2, "SE@43")])
        self.assertEqual([(card["rarity"], card["quantity"]) for card in cards], [("SE", 1), ("SE@43", 3)])
        rows = list(csv.reader(io.StringIO(exports.to_csv_bytes(cards).decode("utf-8-sig"))))
        self.assertEqual(rows[1][1], "시크릿 레어")
        self.assertEqual(rows[2][1], "시크릿 레어（SPECIAL BLUE Ver.）")

    def test_duplicate_copies_are_summed_by_cid(self):
        cards = exports.aggregate([entry(1, "Ash"), entry(2, "Nibiru"), entry(1, "Ash"), entry(1, "Ash", 2)])
        self.assertEqual([(c["cid"], c["quantity"]) for c in cards], [(1, 4), (2, 1)])

    def test_conflicting_names_for_same_cid_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "CID 1"):
            exports.aggregate([entry(1, "Ash"), entry(1, "Ash Blossom")])

    def test_missing_invalid_cid_or_name_is_rejected(self):
        for bad in [entry(None, "Ash"), entry(True, "Ash"), entry(0, "Ash"), entry(-5, "Ash"), entry("12950", "Ash"),
                    entry(1, "  ")]:
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                exports.aggregate([bad])


class FileTest(unittest.TestCase):
    optional_fields = ("name_ko", "name_ja", "name_en", "cid")
    cards = [
        {"cid": 12950, "name": "灰流うらら", "name_ko": "하루 우라라", "name_ja": "灰流うらら",
         "name_en": "Ash Blossom & Joyous Spring", "quantity": 3, "rarity": "SR"},
        {"cid": 99, "name": "=HYPERLINK(\"http://x\")", "name_ko": None, "name_ja": "-ja",
         "name_en": "@en", "quantity": 1, "rarity": "N"},
    ]

    def test_csv_round_trip_with_bom_and_formula_neutralized(self):
        data = exports.to_csv_bytes(self.cards, optional_fields=self.optional_fields)
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["카드명", "레어도", "수량", "한국어 이름", "일본어 이름", "영어 이름", "공식 CID"])
        self.assertEqual(rows[1], ["灰流うらら", "슈퍼 레어", "3", "하루 우라라", "灰流うらら", "Ash Blossom & Joyous Spring", "12950"])
        self.assertEqual(rows[2], ["'=HYPERLINK(\"http://x\")", "노멀", "1", "", "'-ja", "'@en", "99"])

    def test_xlsx_round_trip_keeps_text_never_formula(self):
        sheet = load_workbook(io.BytesIO(exports.to_xlsx_bytes(self.cards, optional_fields=self.optional_fields))).active
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(rows[0], ["카드명", "레어도", "수량", "한국어 이름", "일본어 이름", "영어 이름", "공식 CID"])
        self.assertEqual(rows[1], ["灰流うらら", "슈퍼 레어", 3, "하루 우라라", "灰流うらら", "Ash Blossom & Joyous Spring", 12950])
        self.assertEqual(rows[2][0], "=HYPERLINK(\"http://x\")")
        self.assertEqual(sheet["A3"].data_type, "s")

    def test_default_export_contains_only_display_name_quantity_and_rarity(self):
        csv_rows = list(csv.reader(io.StringIO(exports.to_csv_bytes(self.cards).decode("utf-8-sig"))))
        sheet = load_workbook(io.BytesIO(exports.to_xlsx_bytes(self.cards))).active
        self.assertEqual(csv_rows[0], ["카드명", "레어도", "수량"])
        self.assertEqual(csv_rows[1], ["灰流うらら", "슈퍼 레어", "3"])
        self.assertEqual(list(sheet.values), [
            ("카드명", "레어도", "수량"),
            ("灰流うらら", "슈퍼 레어", 3),
            ('=HYPERLINK("http://x")', "노멀", 1),
        ])
        self.assertEqual(sheet["A3"].data_type, "s")

    def test_all_sixteen_optional_column_combinations_in_both_formats(self):
        labels = ("한국어 이름", "일본어 이름", "영어 이름", "공식 CID")
        values = ("하루 우라라", "灰流うらら", "Ash Blossom & Joyous Spring", 12950)
        for count in range(5):
            for indices in combinations(range(4), count):
                fields = tuple(self.optional_fields[i] for i in reversed(indices))
                with self.subTest(fields=fields):
                    expected_headers = ["카드명", "레어도", "수량", *[labels[i] for i in indices]]
                    expected_values = ["灰流うらら", "슈퍼 레어", 3, *[values[i] for i in indices]]
                    csv_rows = list(csv.reader(io.StringIO(
                        exports.to_csv_bytes(self.cards[:1], optional_fields=fields).decode("utf-8-sig"))))
                    sheet = load_workbook(io.BytesIO(
                        exports.to_xlsx_bytes(self.cards[:1], optional_fields=fields))).active
                    self.assertEqual(csv_rows, [expected_headers, [str(value) for value in expected_values]])
                    self.assertEqual(list(sheet.values), [tuple(expected_headers), tuple(expected_values)])

    def test_unknown_optional_field_is_an_explicit_error(self):
        for writer in (exports.to_csv_bytes, exports.to_xlsx_bytes):
            for cards in ([], self.cards):
                with self.subTest(writer=writer.__name__, empty=not cards), self.assertRaises(ValueError):
                    writer(cards, optional_fields=("name_ko", "typo"))


if __name__ == "__main__":
    unittest.main()
