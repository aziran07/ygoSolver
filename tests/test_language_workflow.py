"""User requirements: per-card edition, mixed decks, explicit uncertainty and errors."""
import copy
import csv
import io
import unittest
from unittest import mock

from openpyxl import load_workbook
from streamlit.delta_generator import DeltaGenerator

import card_language
import exports
import rarities
import test_app as fixtures
import test_app_prices as prices
from price_test_fixtures import capture_downloads, group, product


def detection(locale):
    return {"locale": locale, "status": "classified" if locale else "review",
            "reason": "독립 테스트 문자 근거" if locale else "글자가 흐려 확인 필요",
            "evidence": {"sample": "unchanged"}, "elapsed_seconds": 0.01}


class MixedEditionExportTest(unittest.TestCase):
    def entry(self, locale="ko", quantity=1):
        return {"cid": 12950, "name": "하루 우라라", "name_ko": "하루 우라라", "name_ja": "灰流うらら",
                "name_en": "Ash Blossom & Joyous Spring", "rarity": "N", "locale": locale, "quantity": quantity}

    def test_same_edition_merges_but_korean_and_japanese_remain_separate_in_both_formats(self):
        cards = exports.aggregate([self.entry("ko", 2), self.entry("ja", 4), self.entry("ko", 3)])
        self.assertEqual([(c["locale"], c["quantity"]) for c in cards], [("ko", 5), ("ja", 4)])
        rows = list(csv.DictReader(io.StringIO(exports.to_csv_bytes(cards).decode("utf-8-sig"))))
        self.assertEqual([(r["판본"], r["수량"]) for r in rows], [("한국판", "5"), ("일본판", "4")])
        sheet = load_workbook(io.BytesIO(exports.to_xlsx_bytes(cards))).active
        self.assertEqual(list(sheet.values), [("카드명", "레어도", "판본", "수량"),
                         ("하루 우라라", "노멀", "한국판", 5), ("하루 우라라", "노멀", "일본판", 4)])

    def test_unknown_or_unsupported_edition_blocks_every_export_path(self):
        missing = self.entry()
        del missing["locale"]
        for card in (missing, *[self.entry(locale) for locale in (None, "en", "", "JP")]):
            for operation in (exports.aggregate, exports.to_csv_bytes, exports.to_xlsx_bytes):
                with self.subTest(card=card, operation=operation.__name__), self.assertRaises(ValueError):
                    operation([card])

    def test_price_cannot_belong_to_the_other_physical_edition(self):
        from test_price_exports import quote
        card = {**self.entry("ko"), "price": quote()}
        for operation in (exports.to_csv_bytes, exports.to_xlsx_bytes):
            with self.subTest(operation=operation.__name__), self.assertRaisesRegex(ValueError, "판본"):
                operation([card], include_price=True)


class LanguagePriceAppTest(unittest.TestCase):
    def setUp(self):
        self.fixture = prices.PriceAppTest()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)

    def test_unknown_edition_blocks_lookup_and_download_until_user_selects_it(self):
        unknown = {**prices.row(), "locale": None}
        with mock.patch.object(rarities, "get_rarities", return_value=["N"]) as rarity:
            app = self.fixture.start([unknown])
            rarity.assert_not_called()
        self.fixture.fetch.assert_not_called()
        self.fixture.read.assert_not_called()
        self.fixture.collect.assert_not_called()
        self.assertEqual(fixtures.download_buttons_disabled(app), [True, True])
        self.assertIsNone(app.selectbox(key="locale_one").value)
        app.selectbox(key="locale_one").set_value("ja").run()
        self.assertEqual(app.metric[0].value, "240원")
        self.assertEqual(app.session_state["rows"][0]["locale"], "ja")
        app.checkbox(key="confirmed").check().run()
        self.assertEqual(fixtures.download_buttons_disabled(app), [False, False])

    def test_old_session_never_takes_legacy_global_price_locale_as_physical_evidence(self):
        old = prices.row()
        del old["locale"]
        app = self.fixture.start([old])
        self.assertIsNone(app.session_state["rows"][0]["locale"])
        self.fixture.read.assert_not_called()
        self.assertEqual(fixtures.download_buttons_disabled(app), [True, True])

    def test_mixed_deck_prices_and_actual_downloads_keep_editions_separate(self):
        korean = {**prices.row("ko", 2), "locale": "ko"}
        japanese = {**prices.row("ja", 3), "locale": "ja"}
        def read(number, locale, *args):
            amount = 800 if locale == "ko" else 240
            return {**group(number, [product(number=number, price=amount, locale=locale)]), "locale": locale}
        self.fixture.read.side_effect = read
        files = {}
        with mock.patch.object(DeltaGenerator, "download_button", capture_downloads(files)):
            app = self.fixture.start([korean, japanese])
            self.assertFalse(app.exception)
            table = app.dataframe[0].value
            self.assertEqual(list(zip(table["판본"], table["수량"], table["단가(원)"], table["합계(원)"])),
                             [("한국판", 2, 800, 1600), ("일본판", 3, 240, 720)])
            app.checkbox(key="confirmed").check().run()
        for extension in ("csv", "xlsx"):
            self.assertTrue(files[extension])
        csv_rows = list(csv.DictReader(io.StringIO(files["csv"].decode("utf-8-sig"))))
        self.assertEqual([(r["판본"], r["합계(원)"]) for r in csv_rows], [("한국판", "1600"), ("일본판", "720")])
        sheet = load_workbook(io.BytesIO(files["xlsx"])).active
        self.assertEqual([tuple(row[2:6]) for row in list(sheet.values)[1:]],
                         [("한국판", 2, 800, 1600), ("일본판", 3, 240, 720)])

    def test_manual_override_keeps_ocr_evidence_changes_price_and_invalidates_confirmation(self):
        source = detection("ja")
        row = {**prices.row(), "language_detection": copy.deepcopy(source), "language_error": None}
        app = self.fixture.start([row])
        app.checkbox(key="confirmed").check().run()
        app.selectbox(key="locale_one").set_value("ko").run()
        self.assertFalse(app.exception)
        self.assertFalse(app.checkbox(key="confirmed").value)
        self.assertEqual(app.session_state["rows"][0]["language_detection"], source)
        self.assertEqual(app.metric[0].label, "단가 (한국판)")
        self.assertTrue(any(call.args[1] == "ko" for call in self.fixture.read.call_args_list))
        app.radio[0].set_value("영어").run()
        self.assertEqual(app.session_state["rows"][0]["locale"], "ko")

    def test_edition_change_reloads_local_rarities_and_rejects_unprinted_edition(self):
        def available(cid, *, locale):
            return ["SR", "UR"] if locale == "ko" else ["N", "UR"]
        with mock.patch.object(rarities, "get_rarities", side_effect=available) as lookup:
            app = self.fixture.start()
            self.assertEqual(app.selectbox(key="rarity_one").value, "N")
            app.checkbox(key="confirmed").check().run()
            app.selectbox(key="locale_one").set_value("ko").run()
            self.assertFalse(app.exception)
            self.assertEqual(app.selectbox(key="rarity_one").options, ["슈퍼 레어", "울트라 레어"])
            self.assertEqual(app.selectbox(key="rarity_one").value, "SR")
            self.assertFalse(app.checkbox(key="confirmed").value)
            self.assertEqual({call.kwargs["locale"] for call in lookup.call_args_list}, {"ko", "ja"})
        with mock.patch.object(rarities, "get_rarities", side_effect=rarities.RarityUnavailable("no ko printing")):
            app = self.fixture.start([{**prices.row(), "locale": "ko"}])
            self.assertFalse(app.exception)
            self.assertEqual(fixtures.download_buttons_disabled(app), [True, True])
            self.assertFalse(app.dataframe)


class RecognitionLanguageAppTest(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.UploadTest()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)

    def recognize(self):
        fixtures.FakeRecognizer.outcomes = [["recognized", "recognized"]]
        app = fixtures.start_app([]).run()
        self.fixture.upload(app, ("mixed.png", fixtures.png_bytes("white")))
        self.fixture.recognize_button(app).click().run()
        self.assertFalse(app.exception)
        return app

    def test_each_detected_card_gets_its_own_edition_using_original_image(self):
        with mock.patch.object(card_language, "detect_language", side_effect=[detection("ko"), detection("ja")]) as infer:
            app = self.recognize()
        self.assertEqual([r["locale"] for r in app.session_state["rows"]], ["ko", "ja"])
        self.assertEqual(infer.call_count, 2)
        self.assertEqual(infer.call_args_list[0].args[0].size, (64, 48))

    def test_invalid_detector_contract_fails_recognition_instead_of_becoming_review(self):
        for invalid in (detection("en"), {**detection(None), "status": "classified"},
                        {**detection("ko"), "status": "review"}):
            with self.subTest(result=invalid), mock.patch.object(card_language, "detect_language", return_value=invalid):
                app = self.recognize()
            self.assertEqual(app.session_state["rows"], [])
            self.assertIn("판본 판별 결과가 올바르지 않습니다", app.session_state["recognition_error"])
            self.assertEqual(fixtures.download_buttons_disabled(app), [True, True])

    def test_ocr_runtime_error_is_visible_and_distinct_from_review_with_manual_recovery(self):
        with mock.patch.object(card_language, "detect_language", side_effect=[
                card_language.LanguageDetectionError("Korean OCR model missing"), detection(None)]):
            app = self.recognize()
        failed, uncertain = app.session_state["rows"]
        self.assertIn("Korean OCR model missing", failed["language_error"])
        self.assertIsNone(uncertain["language_error"])
        self.assertEqual([r["locale"] for r in (failed, uncertain)], [None, None])
        self.assertEqual(fixtures.download_buttons_disabled(app), [True, True])
        visible = " ".join(e.value for kind in ("error", "warning", "caption", "markdown") for e in app.get(kind))
        self.assertIn("Korean OCR model missing", visible)
        for row in (failed, uncertain):
            app.selectbox(key=f"locale_{row['key']}").set_value("ko").run()
        app.checkbox(key="confirmed").check().run()
        self.assertEqual(fixtures.download_buttons_disabled(app), [False, False])
        self.assertIn("Korean OCR model missing", app.session_state["rows"][0]["language_error"])
        visible = " ".join(e.value for kind in ("success", "warning", "info") for e in app.get(kind))
        self.assertNotIn("판본 확인 필요 -", visible)


if __name__ == "__main__":
    unittest.main()
