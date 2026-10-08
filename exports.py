"""Deck list aggregation and XLSX / CSV export."""

import csv
import io

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import rarities

# Columns the user can add to the export, in their fixed output order between 수량 and 레어도.
OPTIONAL_COLUMNS = {"name_ko": "한국어 이름", "name_ja": "일본어 이름", "name_en": "영어 이름", "cid": "공식 CID"}

# Spreadsheet apps treat cells starting with these characters as formulas.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def validate_quantity(value):
    """Return value if it is a positive whole int; bool, float and Fraction are rejected."""
    if isinstance(value, bool) or not isinstance(value, int):
        raise ValueError(f"수량은 1 이상의 정수여야 합니다: {value!r}")
    if value < 1:
        raise ValueError(f"수량은 1 이상이어야 합니다: {value}")
    return value


def aggregate(entries):
    """Sum quantities of entries sharing the same official CID and rarity.

    Each entry: {'cid': int, 'rarity': rarity code, 'name': str, 'name_ko', 'name_ja', 'name_en', 'quantity': int}.
    The same CID with different rarities stays on separate lines, but every line of one CID must
    carry the same card name. Returns one dict per (CID, rarity) in first-seen order.
    """
    by_cid_rarity = {}
    name_by_cid = {}
    for entry in entries:
        cid = entry["cid"]
        if isinstance(cid, bool) or not isinstance(cid, int) or cid <= 0:
            raise ValueError(f"올바른 공식 CID(양의 정수)가 없는 항목은 내보낼 수 없습니다: {entry['name']!r} (CID {cid!r})")
        quantity = validate_quantity(entry["quantity"])
        name = entry["name"].strip()
        if not name:
            raise ValueError(f"CID {cid} 항목의 카드명이 비어 있습니다.")
        rarity = entry.get("rarity")
        if rarity not in rarities.RARITY_LABELS:
            raise ValueError(f"CID {cid} 항목의 레어도가 올바르지 않습니다: {rarity!r}. 레어도를 선택하세요.")
        if name_by_cid.setdefault(cid, name) != name:
            raise ValueError(
                f"같은 카드(CID {cid})에 서로 다른 카드명이 입력되었습니다: "
                f"{name_by_cid[cid]!r} / {name!r}. 하나로 맞춰 주세요."
            )
        key = (cid, rarity)
        if key not in by_cid_rarity:
            by_cid_rarity[key] = {**entry, "name": name, "quantity": 0}
        by_cid_rarity[key]["quantity"] += quantity
    return list(by_cid_rarity.values())


def selected_fields(optional_fields):
    """Return the chosen optional fields in canonical order; unknown fields raise ValueError."""
    unknown = set(optional_fields) - OPTIONAL_COLUMNS.keys()
    if unknown:
        raise ValueError(f"알 수 없는 내보내기 열입니다: {sorted(unknown)!r}")
    return [field for field in OPTIONAL_COLUMNS if field in optional_fields]


def headers(optional_fields=()):
    return ["카드명", "수량", *(OPTIONAL_COLUMNS[field] for field in selected_fields(optional_fields)), "레어도"]


def row_values(card, optional_fields=()):
    return [card["name"], card["quantity"], *(card[field] for field in selected_fields(optional_fields)),
            rarities.RARITY_LABELS[card["rarity"]]]


HEADERS = headers(OPTIONAL_COLUMNS)


def _neutralize_formula(value):
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def to_csv_bytes(cards, optional_fields=()):
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(headers(optional_fields))
    for card in cards:
        writer.writerow([_neutralize_formula("" if v is None else v) for v in row_values(card, optional_fields)])
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx_bytes(cards, optional_fields=()):
    header_row = headers(optional_fields)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "덱"
    sheet.append(header_row)
    for card in cards:
        sheet.append(row_values(card, optional_fields))
        # Force text cells so names like "=..." are stored as text, never as formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.column_dimensions["A"].width = 36
    for name_column in ("한국어 이름", "일본어 이름", "영어 이름"):
        if name_column in header_row:
            sheet.column_dimensions[get_column_letter(header_row.index(name_column) + 1)].width = 30
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
